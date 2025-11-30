# generar_informe.py
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import pandas as pd
import os
from datetime import datetime, timedelta
from io import BytesIO

# Solo matplotlib (sin seaborn)
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

# Estilo limpio
plt.rcParams.update({
    'font.size': 10,
    'font.family': 'Times New Roman',
    'axes.edgecolor': 'black',
    'axes.linewidth': 0.8,
    'axes.grid': True,
    'grid.color': 'lightgray',
    'grid.linestyle': '--',
    'grid.alpha': 0.7,
    'figure.autolayout': True
})

LOGO_PATH = "img/logo.png"

def _nombre_mes_es(mes_num):
    meses = [
        'enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio',
        'julio', 'agosto', 'septiembre', 'octubre', 'noviembre', 'diciembre'
    ]
    if 1 <= mes_num <= 12:
        return meses[mes_num - 1].capitalize()
    return ''

def _parse_fecha_str(s):
    if not s:
        return None
    try:
        return datetime.strptime(s, '%d/%m/%Y')
    except ValueError:
        try:
            return pd.to_datetime(s, dayfirst=True, errors='coerce')
        except Exception:
            return None

def apply_font(paragraph, name='Times New Roman', size=12, bold=False, italic=False, color=None):
    for run in paragraph.runs:
        run.font.name = name
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        if color:
            run.font.color.rgb = RGBColor(*color)

def crear_informe(start_date_str: str, end_date_str: str, csv_path: str = 'incidentes.csv') -> bytes:
    dt_start = _parse_fecha_str(start_date_str)
    dt_end = _parse_fecha_str(end_date_str)

    if dt_start is None or dt_end is None:
        dt_end = datetime.now()
        dt_start = dt_end - timedelta(days=7)
        start_date_str = dt_start.strftime('%d/%m/%Y')
        end_date_str = dt_end.strftime('%d/%m/%Y')

    mes_ano_portada = ''
    if dt_start and dt_end:
        if dt_start.month == dt_end.month and dt_start.year == dt_end.year:
            mes_ano_portada = f"{_nombre_mes_es(dt_start.month)} {dt_start.year}"
        else:
            mes_ano_portada = f"{_nombre_mes_es(dt_start.month)} {dt_start.year} – {_nombre_mes_es(dt_end.month)} {dt_end.year}"

    doc = Document()

    # Estilo global
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)

    section = doc.sections[0]
    section.left_margin = section.right_margin = Inches(1)
    section.top_margin = section.bottom_margin = Inches(1)

    # === PORTADA CENTRADA (sin hoja en blanco) ===
    for _ in range(10):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.space_after = Pt(0)
        p.space_before = Pt(0)

    if os.path.exists(LOGO_PATH):
        try:
            logo_para = doc.add_paragraph()
            logo_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run_logo = logo_para.add_run()
            run_logo.add_picture(LOGO_PATH, width=Inches(2.0))
            doc.add_paragraph()
        except Exception as e:
            print(f"Advertencia: no se pudo insertar el logo: {e}")

    titulo = doc.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_titulo = titulo.add_run("INFORME DE AFECTACIONES CRÍTICAS")
    run_titulo.font.size = Pt(24)
    run_titulo.font.bold = True
    run_titulo.font.name = 'Times New Roman'

    doc.add_paragraph()

    periodo = doc.add_paragraph()
    periodo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_periodo = periodo.add_run(f"Periodo Cubierto: {start_date_str} – {end_date_str}")
    run_periodo.font.size = Pt(14)
    run_periodo.font.name = 'Times New Roman'

    if mes_ano_portada:
        doc.add_paragraph()
        mes_par = doc.add_paragraph()
        mes_par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_mes = mes_par.add_run(f"Mes: {mes_ano_portada}")
        run_mes.font.size = Pt(12)
        run_mes.font.name = 'Times New Roman'

    # ÚNICO salto de página: al final de la portada
    doc.add_page_break()

    # === ENCABEZADO Y PIE ===
    header = section.header
    header_para = header.paragraphs[0]
    header_para.text = "INFORME DE AFECTACIONES CRÍTICAS"
    header_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    apply_font(header_para, size=10)

    footer = section.footer
    footer_para = footer.paragraphs[0]
    footer_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run_page = footer_para.add_run()
    fldSimple = OxmlElement('w:fldSimple')
    fldSimple.set(qn('w:instr'), 'PAGE')
    run_page._element.append(fldSimple)
    apply_font(footer_para, size=10)

    # === CUERPO DEL DOCUMENTO ===
    titulo_sec = doc.add_heading("INFORME DE AFECTACIONES CRÍTICAS", level=1)
    apply_font(titulo_sec, size=16, bold=True)

    doc.add_heading("Objetivo", level=2)
    texto_objetivo = doc.add_paragraph(
        "Brindar visibilidad sobre los incidentes que impactan la disponibilidad de los servicios, "
        "destacando el rol y las acciones realizadas por el área de Ciberseguridad en la identificación, "
        "análisis y resolución de estos."
    )
    apply_font(texto_objetivo, size=12)

    # === CARGAR DATOS ===
    datos_incidentes = []
    df_filtrado = pd.DataFrame()

    if os.path.exists(csv_path):
        codificaciones = ['utf-8', 'latin-1', 'iso-8859-1', 'cp1252']
        df = None
        for enc in codificaciones:
            try:
                df = pd.read_csv(csv_path, encoding=enc, sep=',', skiprows=1)
                break
            except:
                continue

        if df is not None and len(df) > 0:
            df.columns = df.columns.str.strip().str.replace('"', '', regex=False)
            columnas_necesarias = ['Nro.caso', 'Servicio', 'Tipo de afectación', 'Fecha inicio', 'Hora inicio', 'Fecha fin', 'Hora fin', 'Min']
            cols_existentes = [c for c in columnas_necesarias if c in df.columns]
            if cols_existentes:
                df_filtrado = df[cols_existentes].copy()
                for col in ['Fecha inicio', 'Fecha fin']:
                    if col in df_filtrado.columns:
                        df_filtrado[col] = pd.to_datetime(df_filtrado[col], dayfirst=True, errors='coerce')

                try:
                    start_dt = pd.to_datetime(start_date_str, dayfirst=True, errors='coerce')
                    end_dt = pd.to_datetime(end_date_str, dayfirst=True, errors='coerce')
                    if pd.notna(start_dt) and pd.notna(end_dt) and 'Fecha inicio' in df_filtrado.columns and 'Fecha fin' in df_filtrado.columns:
                        mask = (df_filtrado['Fecha inicio'] <= end_dt) & (df_filtrado['Fecha fin'] >= start_dt)
                        df_filtrado = df_filtrado[mask]
                except:
                    pass

                if df_filtrado.empty:
                    df_filtrado = df[cols_existentes].copy()

                df_filtrado = df_filtrado[df_filtrado['Servicio'].notna() & (df_filtrado['Servicio'] != '')]

                total_min = 0
                avg_min = 0.0
                if 'Min' in df_filtrado.columns:
                    df_unicos = df_filtrado.drop_duplicates(subset=['Nro.caso'])
                    mins_series = pd.to_numeric(df_unicos['Min'], errors='coerce').dropna()
                    if not mins_series.empty:
                        total_min = int(mins_series.sum())
                        avg_min = float(mins_series.mean())

                def minutos_a_hora(mins):
                    if mins == 0:
                        return "00:00"
                    h = mins // 60
                    m = mins % 60
                    return f"{int(h):02d}:{int(m):02d}"

                total_hm = minutos_a_hora(total_min)
                avg_hm = minutos_a_hora(int(avg_min))
                texto_objetivo.add_run(f"\n\nTotal acumulado: {total_hm} horas | Promedio por incidente: {avg_hm} horas")

                columnas_tabla = ['Nro.caso', 'Servicio', 'Tipo de afectación', 'Fecha inicio', 'Hora inicio', 'Fecha fin', 'Hora fin', 'Min']
                for _, row in df_filtrado.iterrows():
                    fila = []
                    for col in columnas_tabla:
                        val = row[col] if col in row else ""
                        if 'fecha' in col.lower() and pd.notna(val):
                            if isinstance(val, pd.Timestamp):
                                val = val.strftime('%d/%m/%Y')
                        elif 'hora' in col.lower() and pd.notna(val):
                            if isinstance(val, str):
                                parts = val.split(':')
                                if len(parts) >= 2:
                                    val = f"{parts[0].zfill(2)}:{parts[1].zfill(2)}"
                        elif col.lower() == 'min' and pd.notna(val):
                            try:
                                num = pd.to_numeric(val, errors='coerce')
                                if pd.notna(num):
                                    fila.append(str(int(num)))
                                    continue
                            except:
                                pass
                        fila.append(str(val) if pd.notna(val) else "")
                    datos_incidentes.append(fila)

    if not datos_incidentes:
        datos_ejemplo = [
            ["EJ123", "Mi Bancolombia", "Alta", "07/11/2025", "10:04", "07/11/2025", "13:57", "233"],
            ["EJ124", "Portal de Contenidos", "Media", "07/11/2025", "10:04", "07/11/2025", "13:57", "233"],
            ["EJ125", "Nequi Plata", "Alta", "09/11/2025", "12:58", "09/11/2025", "13:50", "52"],
            ["EJ126", "Corresponsales", "Baja", "10/11/2025", "08:30", "10/11/2025", "09:15", "45"],
            ["EJ127", "App Móvil", "Alta", "11/11/2025", "14:20", "11/11/2025", "16:40", "140"],
        ]
        datos_incidentes = datos_ejemplo
        df_filtrado = pd.DataFrame(datos_ejemplo, columns=['Nro.caso', 'Servicio', 'Tipo de afectación', 'Fecha inicio', 'Hora inicio', 'Fecha fin', 'Hora fin', 'Min'])

    # === ALCANCE Y METODOLOGÍA ===
    doc.add_heading("Alcance", level=2)
    puntos_alcance = [
        "Incidentes con alto impacto (afectación parcial o total de servicios críticos).",
        "Incidentes recurrentes o que demandan una participación significativa de Ciber para su solución.",
        "Incidentes donde la causa se relacione directamente con cambios, configuraciones o políticas de seguridad, "
        "infraestructura de seguridad o aplicativos críticos."
    ]
    for punto in puntos_alcance:
        p = doc.add_paragraph(punto, style='List Bullet')
        apply_font(p, size=12)

    doc.add_heading("Metodología de Recolección de Información", level=2)
    metodo_p = doc.add_paragraph()
    apply_font(metodo_p, size=12)
    metodo_p.add_run("• Frecuencia: ").bold = True
    metodo_p.add_run("Se documentan todos los incidentes ocurridos en la semana.\n")
    metodo_p.add_run("• Fuente: ").bold = True
    metodo_p.add_run("Notificaciones oficiales de TI.\n")
    metodo_p.add_run("• Formato: ").bold = True
    metodo_p.add_run("Información estandarizada desde TI.")

    # === GRÁFICAS ===
    if not df_filtrado.empty:
        doc.add_heading("Análisis Visual de Incidentes", level=2)

        # Gráfica 1: Top 5 servicios
        plt.figure(figsize=(6, 4))
        top_servicios = df_filtrado['Servicio'].value_counts().head(5)
        bars = plt.barh(top_servicios.index, top_servicios.values, color='#1F4E78')
        plt.xlabel('Número de Incidentes')
        plt.title('Top 5 Servicios con Más Incidentes')
        plt.gca().invert_yaxis()
        for bar in bars:
            plt.text(bar.get_width() + 0.1, bar.get_y() + bar.get_height()/2, 
                     str(int(bar.get_width())), va='center')
        plt.tight_layout()
        img_buffer1 = BytesIO()
        plt.savefig(img_buffer1, format='png', dpi=150)
        plt.close()
        img_buffer1.seek(0)

        grafica1 = doc.add_paragraph()
        grafica1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run1 = grafica1.add_run()
        run1.add_picture(img_buffer1, width=Inches(6.0))
        doc.add_paragraph()

        # Gráfica 2: Tipos de afectación
        plt.figure(figsize=(5, 5))
        tipo_afectacion = df_filtrado['Tipo de afectación'].value_counts()
        colors = ['#1F4E78', '#3C7CA8', '#6BAED6', '#A1C4E0', '#CCE5F5']
        wedges, texts, autotexts = plt.pie(
            tipo_afectacion.values, 
            labels=tipo_afectacion.index, 
            autopct='%1.1f%%', 
            startangle=90, 
            colors=colors[:len(tipo_afectacion)],
            textprops={'fontsize': 9}
        )
        plt.title('Distribución por Tipo de Afectación', fontsize=11)
        plt.tight_layout()
        img_buffer2 = BytesIO()
        plt.savefig(img_buffer2, format='png', dpi=150)
        plt.close()
        img_buffer2.seek(0)

        grafica2 = doc.add_paragraph()
        grafica2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run2 = grafica2.add_run()
        run2.add_picture(img_buffer2, width=Inches(5.0))
        doc.add_paragraph()

    # === TOP 5 ===
    doc.add_heading("Top 5 Incidentes con Mayor Duración", level=2)
    if not df_filtrado.empty and 'Min' in df_filtrado.columns:
        try:
            df_top = df_filtrado.copy()
            df_top['Min_numeric'] = pd.to_numeric(df_top['Min'], errors='coerce')
            df_top = df_top.dropna(subset=['Min_numeric'])
            df_top = df_top.drop_duplicates(subset=['Nro.caso'])
            df_top = df_top.sort_values('Min_numeric', ascending=False).head(5)

            for idx, (_, row) in enumerate(df_top.iterrows(), 1):
                mins = int(row['Min_numeric'])
                dur = f"{mins // 60:02d}:{mins % 60:02d}"
                p = doc.add_paragraph()
                p.add_run(f"{idx}. ").bold = True
                p.add_run(f"Incidente {row.get('Nro.caso', '')} – {row.get('Servicio', '')} ")
                p.add_run(f"({row.get('Fecha inicio', '')} {row.get('Hora inicio', '')} a {row.get('Fecha fin', '')} {row.get('Hora fin', '')}) – {dur} h")
                apply_font(p, size=11)
        except Exception as e:
            doc.add_paragraph(f"Error en Top 5: {e}", style='Intense Quote')
    else:
        doc.add_paragraph("No hay datos suficientes para mostrar el Top 5.", style='Intense Quote')

    doc.add_paragraph()

    # === TABLA DE INCIDENTES ===
    doc.add_heading("Detalle de Incidentes", level=2)
    tabla = doc.add_table(rows=1, cols=8)
    tabla.style = 'Table Grid'
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER

    widths = [Inches(w) for w in [0.7, 2.2, 1.0, 0.85, 0.75, 0.85, 0.75, 0.6]]
    for i, width in enumerate(widths):
        tabla.columns[i].width = width

    headers = ['Incidente', 'Servicio', 'Tipo de afec.', 'Fecha inicio', 'Hora inicio', 'Fecha fin', 'Hora fin', 'Min']
    hdr_cells = tabla.rows[0].cells
    for i, header in enumerate(headers):
        cell = hdr_cells[i]
        cell.text = header
        for para in cell.paragraphs:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            apply_font(para, size=10, bold=True, color=(255,255,255))
        tcPr = cell._element.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:fill'), '1F4E78')
        tcPr.append(shd)

    for i, fila in enumerate(datos_incidentes):
        row_cells = tabla.add_row().cells
        bg_color = 'FFFFFF' if i % 2 == 0 else 'F5F9FF'
        for j, valor in enumerate(fila):
            cell = row_cells[j]
            cell.text = valor
            tcPr = cell._element.get_or_add_tcPr()
            shd = OxmlElement('w:shd')
            shd.set(qn('w:fill'), bg_color)
            tcPr.append(shd)
            for para in cell.paragraphs:
                if j == 1:
                    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
                else:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                apply_font(para, size=9)

    if datos_incidentes:
        total_min = sum(int(f[7]) for f in datos_incidentes if f[7].isdigit())
        footer_cells = tabla.add_row().cells
        merged = footer_cells[0].merge(footer_cells[5])
        merged.text = f"Total de incidentes: {len(datos_incidentes)}"
        for para in merged.paragraphs:
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            apply_font(para, size=10, bold=True)
        footer_cells[6].text = "Total Min"
        footer_cells[7].text = str(total_min)
        for c in [footer_cells[6], footer_cells[7]]:
            for para in c.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                apply_font(para, size=10, bold=True)
        for c in footer_cells:
            tcPr = c._element.get_or_add_tcPr()
            shd = OxmlElement('w:shd')
            shd.set(qn('w:fill'), 'E2EFFF')
            tcPr.append(shd)

    # === GUARDAR ===
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()